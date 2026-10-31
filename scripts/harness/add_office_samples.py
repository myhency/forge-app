#!/usr/bin/env python3
"""Adds office samples to the harness fixture: an xlsx (two sheets, numbers, a date, a formula nobody calculated),
a fake docx, a real `report.docx` (heading, bold/italic, bullet and numbered list from python-docx list styles, a 2x2 table and a 14-column table) and `deck.pptx` (title slide, a bullet slide with a nested bullet, a picture-less blank slide), and an xlsx outside the session folder, plus a reply in the session "Read the config" that names them
in inline code (the last message of the turn), so the `결과물` card and the links can be checked.

Run it through launch.sh:  FORGE_HARNESS_EXTRA=scripts/harness/add_office_samples.py scripts/harness/launch.sh
"""
import datetime
import json
import os
import sys
from pathlib import Path

import docx
import openpyxl
import pptx
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

import build_fixture as bf

SESSION_ID = "00000000-0000-4000-8000-000000000006"


def make_book(path, title):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "요약"
    sheet.append(["항목", "금액", "날짜"])
    sheet.append(["커피", 4500, datetime.date(2026, 10, 1)])
    sheet.append(["점심", 12000, datetime.date(2026, 10, 2)])
    sheet.append([title, "=SUM(B2:B3)", None])
    data = book.create_sheet("데이터")
    for row in range(1, 700):
        data.append([row, f"user{row}", row * 1.5])
    book.save(path)


def make_docx(path):
    doc = docx.Document()
    doc.add_heading("분기 보고서", level=1)
    paragraph = doc.add_paragraph("이 문단에는 ")
    paragraph.add_run("굵은 글씨").bold = True
    paragraph.add_run("와 ")
    paragraph.add_run("기울인 글씨").italic = True
    paragraph.add_run("가 있습니다.")
    styled = doc.add_paragraph("서식: ")
    styled.add_run("빨간 글씨").font.color.rgb = RGBColor(0xFF, 0, 0)
    styled.add_run(" ")
    styled.add_run("노란 형광펜").font.highlight_color = WD_COLOR_INDEX.YELLOW
    styled.add_run(" ")
    styled.add_run("밑줄").underline = True
    styled.add_run(" ")
    styled.add_run("취소선").font.strike = True
    doc.add_heading("할 일", level=2)
    doc.add_paragraph("첫째 항목", style="List Bullet")
    doc.add_paragraph("둘째 항목", style="List Bullet")
    doc.add_paragraph("번호 하나", style="List Number")
    doc.add_paragraph("번호 둘", style="List Number")
    table = doc.add_table(rows=2, cols=2)
    for row, values in enumerate([["이름", "금액"], ["커피", "4500"]]):
        for column, value in enumerate(values):
            cell = table.cell(row, column)
            cell.text = value
            fill = "1F3864" if row == 0 else ("FFF2CC" if column == 1 else None)
            if fill:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:val"), "clear")
                shading.set(qn("w:color"), "auto")
                shading.set(qn("w:fill"), fill)
                cell._tc.get_or_add_tcPr().append(shading)
            if row == 0:
                run = cell.paragraphs[0].runs[0]
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    wide = doc.add_table(rows=3, cols=14)
    for row in range(3):
        for column in range(14):
            wide.cell(row, column).text = f"열{column + 1}-행{row + 1} 값"
    doc.save(path)


def make_pptx(path):
    deck = pptx.Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[0])
    first.shapes.title.text = "분기 발표"
    first.placeholders[1].text = "2026년 10월"
    second = deck.slides.add_slide(deck.slide_layouts[1])
    second.shapes.title.text = "핵심 내용"
    frame = second.placeholders[1].text_frame
    frame.text = "매출 증가"
    child = frame.add_paragraph()
    child.text = "신규 고객 +20%"
    child.level = 1
    frame.add_paragraph().text = "비용 절감"
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)


def add(claude_home, projects_root):
    cwd = Path(projects_root) / "my-app"
    folder = cwd / "office"
    folder.mkdir(parents=True, exist_ok=True)
    make_book(folder / "report.xlsx", "합계")
    (folder / "memo.docx").write_bytes(b"PK\x03\x04 not really a docx\x00\x00")
    make_docx(folder / "report.docx")
    make_pptx(folder / "deck.pptx")
    outside = Path("/private/tmp/forge-outside-report.xlsx")
    make_book(outside, "밖")

    text = "만들었습니다.\n\n- `{}`\n- `{}`\n- `{}`\n- `{}`\n- `{}`\n".format(
        folder / "report.xlsx", folder / "memo.docx", folder / "report.docx", folder / "deck.pptx", outside
    )
    transcript = Path(claude_home) / "projects" / bf.escape_project_path(cwd) / f"{SESSION_ID}.jsonl"
    line = {"type": "assistant", "sessionId": SESSION_ID, "message": {"role": "assistant", "content": [{"type": "text", "text": text}]}}
    with transcript.open("a") as handle:
        handle.write(json.dumps(line) + "\n")


if __name__ == "__main__":
    run_dir = Path(os.environ.get("FORGE_HARNESS_DIR", "/private/tmp/forge-harness-home"))
    add(run_dir / ".claude", Path(os.path.realpath(bf.DEFAULT_PROJECTS_ROOT)))
    print(f"office samples added under {bf.DEFAULT_PROJECTS_ROOT}/my-app/office", file=sys.stderr)
